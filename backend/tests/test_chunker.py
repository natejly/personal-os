"""Structure-aware chunking, structured extraction, reindex and dedup. Offline.
Run: PERSONAL_OS_DATA_DIR=/tmp/x python backend/tests/test_chunker.py"""
from __future__ import annotations

import io
import os
import sqlite3
import sys
import tempfile
from pathlib import Path
from typing import Any

os.environ.setdefault("PERSONAL_OS_DATA_DIR", tempfile.mkdtemp(prefix="chunktest-"))
os.environ.setdefault("PERSONAL_OS_AUTH_TOKEN", "test-token")
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from fastapi.testclient import TestClient  # noqa: E402

from personal_os import chunker, repos  # noqa: E402
from personal_os.app import AUTH_TOKEN, app, db, documents, embedder  # noqa: E402
from personal_os.chunker import chunk_blocks, split_table, split_text  # noqa: E402
from personal_os.context import _excerpt_header  # noqa: E402
from personal_os.db import Database  # noqa: E402
from personal_os.embed import EmbedError  # noqa: E402
from personal_os.extract_text import extract_structured, extract_text, markdown_blocks  # noqa: E402



async def _offline(*a: Any, **k: Any) -> Any:
    raise EmbedError("offline test")


embedder.fn = _offline  # never reach a real embedding route
client = TestClient(app, headers={"X-Personal-OS-Token": AUTH_TOKEN})
passed = 0


def check(cond: Any, label: str) -> None:
    global passed
    assert cond, label
    passed += 1


# (a) nested headings -> heading paths, never across an H1
md = """# A

""" + ("intro of A. " * 40) + """

## A.1

""" + ("alpha text. " * 20) + """

# B

""" + ("beta text. " * 20)
chunks = chunk_blocks(markdown_blocks(md), title="doc.md")
paths = [c.heading for c in chunks]
check(["A", "A.1"] in paths, "nested heading path")
check(["B"] in paths and all(c.heading[:1] != ["B"] or "alpha" not in c.text for c in chunks), "no chunk crosses an H1")
check(all(c.ctx.startswith("doc.md > ") for c in chunks), "ctx begins with title and path")

# (b) small siblings under one parent merge; big ones do not
small = "# P\n\n## S1\n\none short.\n\n## S2\n\ntwo short.\n"
cs = chunk_blocks(markdown_blocks(small), title="t")
check(len(cs) == 1 and "one short." in cs[0].text and "two short." in cs[0].text, "small siblings merge")
big = "# P\n\n## S1\n\n" + ("x" * 700) + "\n\n## S2\n\n" + ("y" * 700) + "\n"
cs = chunk_blocks(markdown_blocks(big), title="t", max_chars=1200)
check(len(cs) == 2 and cs[0].heading == ["P", "S1"] and cs[1].heading == ["P", "S2"], "big siblings stay apart")
h1s = chunk_blocks(markdown_blocks("# One\n\nshort\n\n# Two\n\nshort\n"), title="t")
check(len(h1s) == 2, "small H1 sections never merge")

# (c) oversized paragraph
para = ". ".join(f"Sentence number {i} of the long paragraph" for i in range(130)) + "."
check(len(para) > 5000, "fixture is long")
pieces = split_text(para, 1200, 120)
check(len(pieces) >= 4 and all(len(p) <= 1200 for p in pieces), "every piece within max")
check(pieces[0][-30:].split()[-1] in pieces[1], "pieces overlap")
cs = chunk_blocks([{"kind": "para", "text": para, "level": 0, "page": None}], max_chars=1200)
check(all(len(c.text) <= 1200 for c in cs) and len(cs) >= 4, "chunk_blocks splits oversized paragraphs")
blob = "z" * 3000
check(all(len(p) <= 1200 for p in split_text(blob, 1200, 120)), "unbreakable text is hard-cut")

# (d) tables split by rows with the header repeated
table = "| h1 | h2 |\n| --- | --- |\n" + "\n".join(f"| row{i} | {'v' * 40} |" for i in range(60))
tp = split_table(table, 600)
check(len(tp) > 2 and all(len(p) <= 600 for p in tp), "table pieces fit")
check(all(p.startswith("| h1 | h2 |\n| --- | --- |") for p in tp), "header and rule repeated")
check(sum(p.count("row") for p in tp) == 60, "no row lost")

# (e) DOCX
import docx  # noqa: E402

d = docx.Document()
d.add_heading("Budget", level=1)
d.add_paragraph("Totals for the year.")
t = d.add_table(rows=2, cols=2)
t.cell(0, 0).text, t.cell(0, 1).text, t.cell(1, 0).text, t.cell(1, 1).text = "Item", "Cost", "Rent", "1200"
buf = io.BytesIO()
d.save(buf)
blocks = extract_structured("b.docx", buf.getvalue(), "")
kinds = [b["kind"] for b in blocks]
check(kinds == ["heading", "para", "table"] and blocks[0]["level"] == 1, "docx heading, paragraph, table in order")
check("| Item | Cost |" in blocks[2]["text"] and "| Rent | 1200 |" in blocks[2]["text"], "docx table as pipe rows")
plain = extract_text("b.docx", buf.getvalue(), "")
check("Totals for the year." in plain and "Rent" in plain and "|" not in plain.split("Totals")[0], "extract_text still plain text")
cs = chunk_blocks(blocks, title="b.docx")
check(cs and cs[0].heading == ["Budget"] and "1200" in cs[0].text, "docx table survives into chunk text")

# (f) PDF via pypdf blank pages
from pypdf import PdfWriter  # noqa: E402

w = PdfWriter()
w.add_blank_page(200, 200)
w.add_blank_page(200, 200)
pb = io.BytesIO()
w.write(pb)
pdf_blocks = extract_structured("x.pdf", pb.getvalue(), "application/pdf")
check([b["page"] for b in pdf_blocks if b["kind"] == "page"] == [1, 2], "pdf page blocks")
check(chunk_blocks(pdf_blocks, title="x.pdf") == [], "blank pdf yields no chunks, no crash")
pdfish = [{"kind": "page", "level": 0, "text": "", "page": 1}, {"kind": "para", "level": 0, "text": "first page text", "page": 1},
          {"kind": "page", "level": 0, "text": "", "page": 2}, {"kind": "para", "level": 0, "text": "y" * 1190, "page": 2}]
pc = chunk_blocks(pdfish, title="x.pdf")
check([c.page for c in pc] == [1, 2], "chunk page is the page of its first paragraph")

# (g) a heading-only term finds the chunk beneath it
up = client.post("/documents", files={"file": ("guide.md", b"# Quarantine procedure\n\nWash your hands often and stay home.\n\n# Other\n\nunrelated words here.\n", "text/markdown")}).json()
hits = documents.search(None, "quarantine")
check(any(h["document_id"] == up["id"] and "Wash your hands" in h["text"] for h in hits), "heading-only query finds the chunk beneath")
check(hits[0]["heading"] == "Quarantine procedure" and "Quarantine" not in hits[0]["text"], "clean excerpt keeps text, heading in its own column")

# (h) migration is idempotent on an old-schema DB
old = Path(tempfile.mkdtemp(prefix="oldschema-"))
con = sqlite3.connect(old / "personal-os.db")
con.executescript("""CREATE TABLE documents (id TEXT PRIMARY KEY, project_id TEXT, name TEXT NOT NULL, mime TEXT NOT NULL DEFAULT '',
  size INTEGER NOT NULL DEFAULT 0, path TEXT NOT NULL, text TEXT NOT NULL DEFAULT '', chunk_count INTEGER NOT NULL DEFAULT 0, created_at REAL NOT NULL);
  CREATE TABLE chunks (id TEXT PRIMARY KEY, document_id TEXT NOT NULL, idx INTEGER NOT NULL, text TEXT NOT NULL);
  INSERT INTO documents VALUES('d1',NULL,'old.txt','text/plain',5,'/nope','# Old\n\nlegacy body text',1,1.0);
  INSERT INTO chunks VALUES('c1','d1',0,'legacy body text');""")
con.commit()
con.close()
Database(old)
odb = Database(old)  # second run: nothing to add, nothing breaks
with odb.connect() as c:
    cc = {r["name"] for r in c.execute("PRAGMA table_info(chunks)")}
    dc = {r["name"] for r in c.execute("PRAGMA table_info(documents)")}
check({"heading", "page"} <= cc and "content_hash" in dc, "columns added once")
od = repos.Documents(odb)
check(od.reindex("d1") == 1, "legacy document reindexes from stored text")
check(od.search(None, "legacy")[0]["heading"] == "Old", "legacy doc gains its heading after reindex")

# (i) reindex keeps ids; duplicates detected
before = documents.get(up["id"])
with db.tx() as c:
    c.execute("DELETE FROM chunks_fts WHERE document_id=?", (up["id"],))
    c.execute("UPDATE chunks SET heading='', page=NULL WHERE document_id=?", (up["id"],))
r = client.post("/documents/reindex", json={"id": up["id"]})
check(r.status_code == 200 and r.json()["chunks"] == 2, "reindex route reports chunks")
after = documents.get(up["id"])
check(after["id"] == before["id"] and after["chunk_count"] == 2, "reindex keeps the document id")
check(documents.search(None, "quarantine")[0]["heading"] == "Quarantine procedure", "reindex restores headings and FTS")
check(client.post("/documents/reindex", json={"id": "nope"}).status_code == 404, "reindex unknown id is 404")
check(client.post("/documents/reindex").status_code == 200, "reindex all with no body")
dup = client.post("/documents", files={"file": ("again.md", b"# Quarantine procedure\n\nWash your hands often and stay home.\n\n# Other\n\nunrelated words here.\n", "text/markdown")}).json()
check(dup.get("duplicate") is True and dup["id"] == up["id"], "identical bytes return duplicate:true")
check(len([d for d in documents.list(None) if d["name"] in ("guide.md", "again.md")]) == 1, "duplicate not stored twice")

# (j) fallback when the chunker raises
orig = chunker.chunk_blocks


def boom(*a: Any, **k: Any) -> Any:
    raise RuntimeError("chunker broke")


chunker.chunk_blocks = boom
fb = documents.create(None, "fallback.txt", "text/plain", 10, "", "plain fallback text about zebras")
chunker.chunk_blocks = orig
check(fb["chunk_count"] == 1 and documents.search(None, "zebras"), "chunk_text fallback used")
check(repos.chunk_text("a\n\nb") == ["a\n\nb"], "chunk_text unchanged")

# (k) excerpt header
check(_excerpt_header({"name": "n.pdf", "idx": 0, "heading": "Intro > Scope", "page": 3}) == "n.pdf — Intro > Scope (p.3)", "header: section and page")
check(_excerpt_header({"name": "n.pdf", "idx": 0, "heading": "", "page": 4}) == "n.pdf (p.4)", "header: page only")
check(_excerpt_header({"name": "n.txt", "idx": 2, "heading": "", "page": None}) == "n.txt (chunk 3)", "header: fallback")
pv = client.post("/context/preview", json={"query": "quarantine"}).json()
ex = [c for c in pv["chunks"] if c["document_id"] == up["id"]]
check(ex and ex[0]["heading"] == "Quarantine procedure" and "### guide.md — Quarantine procedure" in pv["system_prompt"], "context excerpt shows the section")

print(f"test_chunker: {passed} checks passed")
