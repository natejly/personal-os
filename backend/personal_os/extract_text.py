"""Extract plain text from uploaded files.

Every file is accepted. PDF and Word yield their text. Anything that decodes as
text is kept. A binary file is stored with a short note so the chat can still
find it by name.
"""
from __future__ import annotations

import contextvars
import io
import re
import unicodedata
from collections.abc import Iterable
from pathlib import Path
from typing import Any

from .limits import MAX_UNZIPPED_BYTES, MAX_UPLOAD_BYTES  # noqa: F401  (re-exported: app and tests read them here)

MAX_INDEX_CHARS = 400_000
MAX_PDF_PAGES = 80
# MAX_UNZIPPED_BYTES (limits.py): a zip's declared uncompressed size, summed. A 180 KB docx can inflate to hundreds
# of MB of XML. It scales with the upload cap so a legitimately large docx is still read.
_TRUNCATED = "\n\n[Extract truncated. The full file is stored.]"

TEXT_EXT = {
    ".txt", ".md", ".markdown", ".mdx", ".csv", ".tsv", ".json", ".jsonl", ".ndjson",
    ".yaml", ".yml", ".xml", ".svg", ".html", ".htm", ".css", ".log", ".rst", ".toml",
    ".ini", ".cfg", ".conf", ".env", ".py", ".ts", ".tsx", ".js", ".jsx", ".mjs", ".cjs",
    ".rb", ".go", ".rs", ".java", ".c", ".h", ".cpp", ".hpp", ".cs", ".php", ".swift",
    ".kt", ".scala", ".lua", ".pl", ".r", ".sql", ".sh", ".bash", ".zsh", ".fish",
    ".tex", ".org", ".ipynb", ".rtf", ".ics", ".vcf", ".srt", ".vtt", ".vue", ".svelte",
    ".gradle", ".lock", ".geojson", ".kml", ".gpx", ".m", ".mm",
}
TEXT_MIME = {"application/json", "application/xml", "application/javascript", "application/x-javascript", "application/yaml"}


def safe_upload_name(name: str) -> str:
    """A single path segment. Directory pieces, controls, and bidi marks are dropped."""
    raw = (name or "").replace("\\", "/").replace("\x00", "")
    base = []
    for ch in Path(raw).name:
        if ch in "/\\" or unicodedata.category(ch) == "Cf" or not ch.isprintable():
            continue
        base.append(ch)
    cleaned = "".join(base).strip()
    if cleaned in {"", ".", ".."}:
        return "untitled"
    return cleaned[:180]


def for_index(text: str) -> str:
    """What search stores. The original file is kept whole on disk."""
    if len(text) <= MAX_INDEX_CHARS:
        return text
    return text[:MAX_INDEX_CHARS].rstrip() + _TRUNCATED


Block = dict  # {kind: 'heading'|'para'|'table'|'page', level: int, text: str, page: int | None}


def _blk(kind: str, text: str = "", level: int = 0, page: int | None = None) -> Block:
    return {"kind": kind, "level": level, "text": text, "page": page}


def _md_table(rows: list[list[str]]) -> str:
    rows = [[(c or "").replace("\n", " ").replace("|", "/").strip() for c in r] for r in rows if r]
    if not rows:
        return ""
    w = max(len(r) for r in rows)
    rows = [r + [""] * (w - len(r)) for r in rows]
    lines = ["| " + " | ".join(r) + " |" for r in rows]
    lines.insert(1, "| " + " | ".join("---" for _ in range(w)) + " |")
    return "\n".join(lines)


def markdown_blocks(text: str, page: int | None = None) -> list[Block]:
    """ATX headings, fenced code kept whole, pipe tables, blank-line paragraphs."""
    out: list[Block] = []
    para: list[str] = []
    fence: list[str] | None = None

    def flush() -> None:
        if para:
            body = "\n".join(para).strip()
            if body:
                is_table = all(ln.lstrip().startswith("|") for ln in para) and len(para) >= 2
                out.append(_blk("table" if is_table else "para", body, page=page))
            para.clear()

    for line in text.replace("\r\n", "\n").replace("\r", "\n").split("\n"):
        if fence is not None:
            fence.append(line)
            if line.strip().startswith("```"):
                out.append(_blk("para", "\n".join(fence), page=page))
                fence = None
            continue
        if line.strip().startswith("```"):
            flush()
            fence = [line]
            continue
        m = re.match(r"^(#{1,6})\s+(.+?)\s*#*\s*$", line)
        if m:
            flush()
            out.append(_blk("heading", m.group(2).strip(), level=len(m.group(1)), page=page))
        elif not line.strip():
            flush()
        else:
            para.append(line)
    if fence is not None:  # unterminated fence: keep what we have
        out.append(_blk("para", "\n".join(fence), page=page))
    flush()
    return out


def _guess_heading(para: str) -> int:
    """Conservative PDF heading guess: one short line, no sentence punctuation. 0 = not a heading."""
    line = para.strip()
    if "\n" in line or not 3 <= len(line) <= 80 or line[-1] in ".,;:!?" or not re.search(r"[A-Za-z]", line):
        return 0
    words = line.split()
    if line.isupper():
        return 1
    if 2 <= len(words) <= 6 and all(w[0].isupper() for w in words if len(w) > 3) and sum(1 for w in words if w[0].isupper()) >= 2:
        return 2
    return 0


def extract_structured(name: str, data: bytes, mime: str = "") -> list[Block]:
    """Blocks in reading order: headings, paragraphs, tables and (PDF) page markers."""
    ext = Path(name).suffix.lower()
    if ext == ".pdf" or mime == "application/pdf":
        texts, ocr = _pdf_page_texts(data)
        out: list[Block] = [_blk("para", _OCR_NOTE)] if ocr else []
        for n, page_text in enumerate(texts, start=1):
            out.append(_blk("page", page=n))
            for para in re.split(r"\n\s*\n", page_text):
                para = para.strip()
                if not para:
                    continue
                lvl = _guess_heading(para)
                out.append(_blk("heading", para, level=lvl, page=n) if lvl else _blk("para", para, page=n))
        return out
    if ext == ".docx":
        import docx
        from docx.table import Table
        from docx.text.paragraph import Paragraph

        _check_zip(data)
        d = docx.Document(io.BytesIO(data))
        out = []
        for child in d.element.body.iterchildren():
            tag = child.tag.rsplit("}", 1)[-1]
            if tag == "p":
                para = Paragraph(child, d)
                if not para.text.strip():
                    continue
                style = (para.style.name if para.style is not None else "") or ""
                m = re.match(r"^Heading (\d)$", style)
                if m or style == "Title":
                    out.append(_blk("heading", para.text.strip(), level=int(m.group(1)) if m else 1))
                else:
                    out.append(_blk("para", para.text.strip()))
            elif tag == "tbl":
                t = _md_table([[c.text for c in r.cells] for r in Table(child, d).rows])
                if t:
                    out.append(_blk("table", t))
        return out
    if ext in (".csv", ".tsv"):
        return markdown_blocks(_cap_delimited(data.decode("utf-8", errors="replace")))
    if ext in TEXT_EXT or mime.startswith("text/"):
        return markdown_blocks(data.decode("utf-8", errors="replace"))
    return markdown_blocks(extract_text(name, data, mime))


# What is stored for a file nothing could be read out of; _image_text has its own longer wording.
NO_TEXT_NOTE = "[File {name} ({kind}, {size} bytes). No text could be extracted. The file is stored and can be referred to by name.]"
# Either wording, and nothing else: a document that merely quotes the sentence is real text.
_NO_TEXT_RE = re.compile(
    r"\[File .+? \(.*?, \d+ bytes\)\. No text could be extracted[.;].*?The file is stored and can be referred to by name\.\]",
    re.DOTALL,
)


def has_readable_text(text: str) -> bool:
    """False when nothing came out of a file: an empty extraction (a scan with no text layer and no
    OCR) or one that is only the no-text marker. The row is still stored; the upload route reports it."""
    t = (text or "").strip()
    return bool(t) and _NO_TEXT_RE.fullmatch(t) is None


def extract_text(name: str, data: bytes, mime: str = "") -> str:
    ext = Path(name).suffix.lower()
    mime = (mime or "").split(";")[0].strip().lower()
    if ext == ".docx" or mime == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
        # Tables included (the old paragraph-only join dropped them from read_document too).
        try:
            return "\n\n".join(b["text"] for b in extract_structured(name, data, mime) if b["kind"] != "page")
        except Exception:  # noqa: BLE001 - a damaged file falls through to the plain readers below
            pass
    mime = (mime or "").split(";")[0].strip().lower()
    parsed = _parsed(ext, mime, data)
    if parsed is None:
        parsed = _more_parsed(name, ext, mime, data)
    if parsed is not None:
        return parsed
    if _looks_textual(data, mime, ext):
        return data.decode("utf-8", errors="replace")
    kind = mime or ext or "binary"
    return NO_TEXT_NOTE.format(name=name, kind=kind, size=len(data))


_OCR_NOTE = "[OCR text: this PDF has no text layer, so the words below were read from page images and may contain mistakes]"
# Set only inside extract_both: one upload asks for text and blocks, and a scanned PDF must not be OCRed twice.
_PDF_MEMO: contextvars.ContextVar[dict[str, Any] | None] = contextvars.ContextVar("_PDF_MEMO", default=None)


def extract_both(name: str, data: bytes, mime: str = "") -> tuple[str, list[Block] | None]:
    """extract_text and extract_structured over one PDF parse (and one OCR run). Blocks are None when they fail."""
    token = _PDF_MEMO.set({})
    try:
        text = extract_text(name, data, mime)
        try:
            blocks: list[Block] | None = extract_structured(name, data, mime)
        except Exception:  # noqa: BLE001 - the chunker falls back to the plain text
            blocks = None
        return text, blocks
    finally:
        _PDF_MEMO.reset(token)


def _parsed(ext: str, mime: str, data: bytes) -> str | None:
    try:
        if ext == ".pdf" or mime == "application/pdf":
            texts, ocr = _pdf_page_texts(data)
            body = _bounded(texts, MAX_PDF_PAGES)
            if ocr:
                return _OCR_NOTE + "\n\n" + body
            return body
        if ext == ".docx" or mime == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
            import docx

            _check_zip(data)
            d = docx.Document(io.BytesIO(data))
            parts = (p.text for p in d.paragraphs if p.text.strip())
            return _bounded(parts, MAX_PDF_PAGES)
    except Exception:
        return None
    return None


def _check_zip(data: bytes) -> None:
    """Refuse a zip-based document that is not a zip or would expand past the cap. The caller
    treats the raise like any unreadable format: the file is stored with a note instead."""
    import zipfile

    with zipfile.ZipFile(io.BytesIO(data)) as z:
        if sum(i.file_size for i in z.infolist()) > MAX_UNZIPPED_BYTES:
            raise ValueError("Archive expands too large to extract")


def _bounded(parts: Iterable[str], limit: int) -> str:
    out: list[str] = []
    n = 0
    cut = False
    for i, part in enumerate(parts):
        if i >= limit or n >= MAX_INDEX_CHARS:
            cut = True
            break
        room = MAX_INDEX_CHARS - n
        if len(part) > room:
            out.append(part[:room])
            cut = True
            break
        out.append(part)
        n += len(part) + (2 if out else 0)
    text = "\n\n".join(out)
    if cut:
        return text.rstrip() + _TRUNCATED
    return text


def _looks_textual(data: bytes, mime: str, ext: str) -> bool:
    if mime.startswith("text/") or mime in TEXT_MIME or ext in TEXT_EXT:
        return True
    sample = data[:8192]
    if not sample or b"\x00" in sample:
        return False
    try:
        text = sample.decode("utf-8")
    except UnicodeDecodeError:
        return False
    printable = sum(ch.isprintable() or ch in "\n\r\t" for ch in text)
    return printable / len(text) > 0.9


# ---- more formats (spreadsheets, slides, scans, images, open documents) ----
# Every reader here is optional-import or stdlib, bounded, and returns None on any failure so extract_text degrades to
# its "no text could be extracted" marker instead of raising.
IMAGE_EXT = {".png", ".jpg", ".jpeg", ".webp", ".tif", ".tiff", ".bmp", ".gif"}
SHEET_ROWS, SHEET_COLS = 200, 30        # per sheet; the output says what was cut
DELIMITED_LINES = 2_000                 # csv/tsv kept as text, cut here
MAX_ZIP_PART = 40 * 1024 * 1024         # one XML part inside an office zip; a bomb is refused, not inflated
OCR_PDF_PAGES, OCR_PDF_DPI = 15, 150
OCR_PAGE_TIMEOUT_S, OCR_TOTAL_TIMEOUT_S, OCR_IMAGE_TIMEOUT_S = 20, 120, 30
SCAN_CHARS_PER_PAGE = 12                # fewer characters than this per page on average = no text layer


def _which(binary: str) -> str | None:
    import shutil

    return shutil.which(binary)


def _run(cmd: list[str], timeout: float, stdin: bytes | None = None) -> bytes:
    """Run an external binary, bounded. Raises on a non-zero exit or a timeout. A seam for tests."""
    import subprocess

    done = subprocess.run(cmd, input=stdin, capture_output=True, timeout=timeout, check=True)
    return done.stdout


def _ln(el: Any) -> str:
    return str(el.tag).rsplit("}", 1)[-1]


def _zip_part(z: Any, name: str) -> bytes | None:
    try:
        info = z.getinfo(name)
    except KeyError:
        return None
    if info.file_size > MAX_ZIP_PART:
        return None
    return z.read(name)


def _col_index(ref: str) -> int:
    n = 0
    for ch in ref:
        if not ch.isalpha():
            break
        n = n * 26 + (ord(ch.upper()) - 64)
    return max(0, n - 1)


def _fmt_cell(v: Any) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer() and abs(v) < 1e15:
        return str(int(v))
    return str(v)


def _render_sheet(name: str, dims: str, rows: Iterable[list[str]]) -> str:
    """One sheet as a heading plus a markdown table, cut at SHEET_ROWS x SHEET_COLS with a note saying so."""
    kept: list[list[str]] = []
    total = 0
    widest = 0
    for r in rows:
        if not any(c.strip() for c in r):
            continue
        total += 1
        widest = max(widest, len(r))
        if len(kept) < SHEET_ROWS:
            kept.append(r[:SHEET_COLS])
    head = f"## Sheet: {name}" + (f" ({dims})" if dims else "")
    if not kept:
        return head + "\n\n(empty)"
    notes = []
    if total > len(kept):
        notes.append(f"showing the first {len(kept)} of {total} non-empty rows")
    if widest > SHEET_COLS:
        notes.append(f"showing the first {SHEET_COLS} of {widest} columns")
    body = _md_table(kept)
    return head + "\n\n" + body + (f"\n\n[{'; '.join(notes)}]" if notes else "")


def _xlsx_openpyxl(data: bytes) -> str | None:
    try:
        import openpyxl
    except ImportError:
        return None
    _check_zip(data)  # openpyxl inflates parts whole; the caller falls to the per-part-bounded zip parser
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        out = []
        for ws in wb.worksheets:
            try:
                dims = ws.calculate_dimension()
            except Exception:  # noqa: BLE001 - read-only sheets with no dimension record
                dims = ""
            out.append(_render_sheet(ws.title, dims, ([_fmt_cell(c) for c in row] for row in ws.iter_rows(values_only=True))))
        return "\n\n".join(out)
    finally:
        wb.close()


def _xlsx_stdlib(data: bytes) -> str | None:
    import xml.etree.ElementTree as ET
    import zipfile

    z = zipfile.ZipFile(io.BytesIO(data))
    wb = _zip_part(z, "xl/workbook.xml")
    if wb is None:
        return None
    rels: dict[str, str] = {}
    rel_xml = _zip_part(z, "xl/_rels/workbook.xml.rels")
    if rel_xml:
        for r in ET.fromstring(rel_xml):
            rels[r.get("Id", "")] = r.get("Target", "")
    shared: list[str] = []
    ss = _zip_part(z, "xl/sharedStrings.xml")
    if ss:
        for si in ET.fromstring(ss):
            shared.append("".join(t.text or "" for t in si.iter() if _ln(t) == "t"))
    sheets: list[tuple[str, str]] = []
    for i, s in enumerate(e for e in ET.fromstring(wb).iter() if _ln(e) == "sheet"):
        rid = next((v for k, v in s.attrib.items() if k.endswith("}id")), "")
        target = rels.get(rid) or f"worksheets/sheet{i + 1}.xml"
        path = target.lstrip("/") if target.startswith("/") else "xl/" + target
        sheets.append((s.get("name", f"Sheet{i + 1}"), path))
    out = []
    for name, path in sheets:
        part = _zip_part(z, path)
        if part is None:
            continue
        dims = ""
        rows: list[list[str]] = []
        for _, el in ET.iterparse(io.BytesIO(part), events=("end",)):
            tag = _ln(el)
            if tag == "dimension":
                dims = el.get("ref", "")
            elif tag == "row":
                cells: list[str] = []
                for c in el:
                    if _ln(c) != "c":
                        continue
                    idx = _col_index(c.get("r", ""))
                    kind = c.get("t", "n")
                    v = next((x for x in c if _ln(x) == "v"), None)
                    if kind == "inlineStr":
                        text = "".join(t.text or "" for t in c.iter() if _ln(t) == "t")
                    elif v is None or v.text is None:
                        text = ""
                    elif kind == "s":
                        text = shared[int(v.text)] if v.text.isdigit() and int(v.text) < len(shared) else ""
                    elif kind == "b":
                        text = "TRUE" if v.text == "1" else "FALSE"
                    else:  # n, str (a formula's cached text), e: the cached value is what the sheet shows
                        text = v.text
                        if kind == "n":
                            try:
                                text = _fmt_cell(float(text))
                            except ValueError:
                                pass
                    if idx > 2000:  # a stray cell far to the right must not allocate a huge row
                        continue
                    cells.extend([""] * (idx - len(cells)))
                    if idx < len(cells):
                        cells[idx] = text
                    else:
                        cells.append(text)
                rows.append(cells)
                el.clear()
        out.append(_render_sheet(name, dims, rows))
    return "\n\n".join(out) if out else None


def _xlsx_text(data: bytes) -> str | None:
    try:
        got = _xlsx_openpyxl(data)
        if got is not None:
            return got
    except Exception:  # noqa: BLE001 - fall to the zip parser, which does not mind a file openpyxl chokes on
        pass
    return _xlsx_stdlib(data)


def _pptx_text(data: bytes) -> str | None:
    import xml.etree.ElementTree as ET
    import zipfile

    z = zipfile.ZipFile(io.BytesIO(data))
    slides = sorted((n for n in z.namelist() if re.fullmatch(r"ppt/slides/slide\d+\.xml", n)),
                    key=lambda n: int(re.search(r"(\d+)\.xml$", n).group(1)))  # type: ignore[union-attr]
    if not slides:
        return None

    def paras(sp: Any) -> list[tuple[int, str]]:
        res = []
        for p in sp.iter():
            if _ln(p) != "p":
                continue
            lvl = 0
            text = []
            for x in p.iter():
                t = _ln(x)
                if t == "pPr":
                    lvl = int(x.get("lvl") or 0)
                elif t == "t":
                    text.append(x.text or "")
                elif t == "br":
                    text.append("\n")
            s = "".join(text).strip()
            if s:
                res.append((lvl, s))
        return res

    def placeholder(sp: Any) -> str:
        for x in sp.iter():
            if _ln(x) == "ph":
                return x.get("type", "body")
        return ""

    out = []
    for n, name in enumerate(slides, start=1):
        root = ET.fromstring(_zip_part(z, name) or b"<x/>")
        title = ""
        body: list[str] = []
        for el in root.iter():
            tag = _ln(el)
            if tag == "sp":
                ph = placeholder(el)
                ps = paras(el)
                if ph in ("title", "ctrTitle") and not title:
                    title = " ".join(t for _, t in ps)
                elif ph not in ("sldNum", "dt", "ftr"):
                    body.extend("  " * lvl + "- " + t.replace("\n", " ") for lvl, t in ps)
            elif tag == "tbl":
                rows = [["".join(t.text or "" for t in c.iter() if _ln(t) == "t") for c in r if _ln(c) == "tc"] for r in el if _ln(r) == "tr"]
                tb = _md_table(rows)
                if tb:
                    body.append(tb)
        notes = ""
        rel = _zip_part(z, f"ppt/slides/_rels/{name.rsplit('/', 1)[-1]}.rels")
        if rel:
            m = re.search(r'Target="[^"]*?(notesSlide\d+\.xml)"', rel.decode("utf-8", errors="replace"))
            nxml = _zip_part(z, f"ppt/notesSlides/{m.group(1)}") if m else None
            if nxml:
                nroot = ET.fromstring(nxml)
                notes = "\n".join(t for el in nroot.iter() if _ln(el) == "sp" and placeholder(el) == "body" for _, t in paras(el))
        block = [f"## Slide {n}" + (f": {title}" if title else "")]
        block.extend(body)
        if notes.strip():
            block.append("Notes: " + notes.strip())
        out.append("\n\n".join([block[0], "\n".join(block[1:])]) if len(block) > 1 else block[0])
    return "\n\n".join(out)


def _odf_text(ext: str, data: bytes) -> str | None:
    import xml.etree.ElementTree as ET
    import zipfile

    z = zipfile.ZipFile(io.BytesIO(data))
    xml = _zip_part(z, "content.xml")
    if xml is None:
        return None
    root = ET.fromstring(xml)
    if ext == ".ods":
        out = []
        for tbl in (e for e in root.iter() if _ln(e) == "table"):
            rows = []
            for r in (e for e in tbl.iter() if _ln(e) == "table-row"):
                rows.append(["".join(c.itertext()) for c in r if _ln(c) == "table-cell"])
            out.append(_render_sheet(next((v for k, v in tbl.attrib.items() if k.endswith("}name")), "Sheet"), "", rows))
        return "\n\n".join(out) or None
    out = []
    for e in root.iter():
        tag = _ln(e)
        if tag in ("h", "p"):
            text = "".join(e.itertext()).strip()
            if text:
                out.append(("## " if tag == "h" else "") + text)
    return "\n\n".join(out) or None


def _rtf_text(data: bytes) -> str | None:
    s = data.decode("latin-1", errors="replace")
    if not s.lstrip().startswith("{\\rtf"):
        return None
    s = re.sub(r"\{\\\*[^{}]*\}", "", s)  # destinations such as {\*\generator ...}
    s = re.sub(r"\{\\(?:fonttbl|colortbl|stylesheet|info)(?:[^{}]|\{[^{}]*\})*\}", "", s)
    s = re.sub(r"\\par[d]?\b ?", "\n", s)
    s = re.sub(r"\\'([0-9a-fA-F]{2})", lambda m: bytes.fromhex(m.group(1)).decode("cp1252", errors="replace"), s)
    s = re.sub(r"\\[a-zA-Z]+-?\d* ?", "", s)
    s = re.sub(r"\\([\\{}])", r"\1", s).replace("{", "").replace("}", "")
    return re.sub(r"\n{3,}", "\n\n", s).strip() or None


def _epub_text(data: bytes) -> str | None:
    import html
    import zipfile

    z = zipfile.ZipFile(io.BytesIO(data))
    names = sorted(n for n in z.namelist() if n.lower().endswith((".xhtml", ".html", ".htm")))

    def chapters() -> Iterable[str]:
        for n in names:
            raw = (_zip_part(z, n) or b"").decode("utf-8", errors="replace")
            raw = re.sub(r"(?is)<(script|style)\b.*?</\1>", "", raw)
            raw = re.sub(r"(?i)</(p|div|h\d|li|tr|br)\s*>|<br\s*/?>", "\n", raw)
            text = html.unescape(re.sub(r"<[^>]+>", "", raw))
            yield re.sub(r"\n\s*\n+", "\n\n", text).strip()

    return _bounded((c for c in chapters() if c), MAX_PDF_PAGES * 4) or None


def _cap_delimited(text: str) -> str:
    lines = text.splitlines()
    if len(lines) <= DELIMITED_LINES and len(text) <= MAX_INDEX_CHARS:
        return text
    kept = "\n".join(lines[:DELIMITED_LINES])[:MAX_INDEX_CHARS]
    return kept + f"\n\n[Showing the first {kept.count(chr(10)) + 1} of {len(lines)} lines. The full file is stored.]"


def _image_text(name: str, ext: str, mime: str, data: bytes) -> str:
    """OCR text of a picture through tesseract when it is installed, else a marker that points at view_image."""
    kind = mime or ext
    if _which("tesseract"):
        try:
            png = data
            try:
                from PIL import Image

                buf = io.BytesIO()
                Image.open(io.BytesIO(data)).convert("RGB").save(buf, "PNG")  # one format tesseract always reads
                png = buf.getvalue()
            except Exception:  # noqa: BLE001 - hand tesseract the original bytes
                pass
            text = _run(["tesseract", "stdin", "stdout"], OCR_IMAGE_TIMEOUT_S, png).decode("utf-8", errors="replace").strip()
            if text:
                return "[OCR text read from the image " + name + "]\n\n" + text
        except Exception:  # noqa: BLE001 - OCR failing is the same as no OCR
            pass
    return (f"[File {name} ({kind}, {len(data)} bytes). No text could be extracted; this is an image, and view_image looks at pictures. "
            "The file is stored and can be referred to by name.]")


def _ocr_pdf(data: bytes, pages: int) -> list[str] | None:
    """OCR the first OCR_PDF_PAGES pages of a PDF with no text layer (pdftoppm + tesseract), or None when either binary is
    missing or nothing came out. Bounded per page and in total."""
    if not (_which("pdftoppm") and _which("tesseract")):
        return None
    import tempfile
    import time

    t0 = time.monotonic()
    out: list[str] = []
    with tempfile.TemporaryDirectory(prefix="grain-ocr-") as d:
        src = Path(d) / "in.pdf"
        src.write_bytes(data)
        for n in range(1, min(pages, OCR_PDF_PAGES) + 1):
            left = OCR_TOTAL_TIMEOUT_S - (time.monotonic() - t0)
            if left <= 1:
                out.append("[OCR stopped: time budget used]")
                break
            base = str(Path(d) / f"p{n}")
            try:
                _run(["pdftoppm", "-r", str(OCR_PDF_DPI), "-f", str(n), "-l", str(n), "-png", "-singlefile", str(src), base], min(OCR_PAGE_TIMEOUT_S, left))
                png = Path(base + ".png").read_bytes()
                out.append(_run(["tesseract", "stdin", "stdout"], min(OCR_PAGE_TIMEOUT_S, max(1.0, left)), png).decode("utf-8", errors="replace").strip())
            except Exception:  # noqa: BLE001 - one bad page does not lose the rest
                out.append("")
    return out if any(out) else None


def _pdf_page_texts(data: bytes) -> tuple[list[str], bool]:
    """(page texts, ocr). A scan with no text layer is OCRed when the binaries exist."""
    memo = _PDF_MEMO.get()
    if memo is not None and "pdf" in memo:
        return memo["pdf"]
    got = _pdf_page_texts_uncached(data)
    if memo is not None:
        memo["pdf"] = got
    return got


def _pdf_page_texts_uncached(data: bytes) -> tuple[list[str], bool]:
    from pypdf import PdfReader

    pages = list(PdfReader(io.BytesIO(data)).pages)[:MAX_PDF_PAGES]
    texts = [(p.extract_text() or "") for p in pages]
    if pages and sum(len(t.strip()) for t in texts) < SCAN_CHARS_PER_PAGE * len(pages):
        got = _ocr_pdf(data, len(pages))
        if got:
            return got, True
    return texts, False


def _more_parsed(name: str, ext: str, mime: str, data: bytes) -> str | None:
    """The formats extract_text does not read inline. None = not mine, or it failed (the caller degrades)."""
    try:
        if ext in (".xlsx", ".xlsm") or mime == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet":
            return _xlsx_text(data)
        if ext == ".pptx" or mime == "application/vnd.openxmlformats-officedocument.presentationml.presentation":
            return _pptx_text(data)
        if ext in (".odt", ".ods", ".odp"):
            return _odf_text(ext, data)
        if ext == ".epub":
            return _epub_text(data)
        if ext == ".rtf":
            return _rtf_text(data)
        if ext in (".csv", ".tsv"):
            return _cap_delimited(data.decode("utf-8", errors="replace"))
        if ext in IMAGE_EXT or (mime.startswith("image/") and mime != "image/svg+xml"):
            return _image_text(name, ext, mime, data)
    except Exception:  # noqa: BLE001 - extraction never raises; a damaged file becomes the "no text" marker
        return None
    return None
